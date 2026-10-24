"""Excel report for the Section 6.3 independent check, Automatic (AI) path.

Reads evaluation/independent-check/ai-results/results.json (from
run_independent_check_ai.mjs) and the team's diagnoses.json, and compares
each check with the deterministic (keyword) path's committed results.

Usage:
  python3 evaluation/build_independent_check_ai_report.py --out "Calder AI Independent Check.xlsx"

Requires openpyxl.
"""

from __future__ import annotations

import argparse
import json
import pathlib
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parents[1]
RESULTS = ROOT / "evaluation" / "independent-check" / "ai-results"
HEADER_FILL = PatternFill("solid", fgColor="1F3A4D")
HEADER_FONT = Font(bold=True, color="FFFFFF")
OK_FILL = PatternFill("solid", fgColor="DCEFE3")
MISS_FILL = PatternFill("solid", fgColor="F8DADA")
EDIT_LABELS = {
    "remove-protection": "Remove a protection",
    "change-liability-cap": "Change a liability cap",
    "shorten-notice-period": "Shorten a notice period",
    "insert-unusual-clause": "Insert an unusual clause",
}


def yes_no(value):
    return None if value is None else ("Yes" if value else "No")


def sheet(workbook, title, headers, rows, widths, wrap=()):
    ws = workbook.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in rows:
        ws.append(row)
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for column in wrap:
        for cell in ws[get_column_letter(column)][1:]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", required=True, type=pathlib.Path)
    args = parser.parse_args()

    run = json.loads((RESULTS / "results.json").read_text(encoding="utf-8"))
    diagnoses = json.loads((RESULTS / "diagnoses.json").read_text(encoding="utf-8"))["diagnoses"]
    rows = run["rows"]
    workbook = Workbook()
    workbook.remove(workbook.active)

    by_type = defaultdict(lambda: dict(edits=0, ai=0, kw=0))
    for r in rows:
        t = by_type[r["editType"]]
        t["edits"] += 1
        t["ai"] += r["beforeCorrect"] + r["afterCorrect"]
        t["kw"] += bool(r["keywordBeforeCorrect"]) + bool(r["keywordAfterCorrect"])
    summary = [
        [EDIT_LABELS.get(k, k), v["edits"], f'{v["ai"]} of {v["edits"] * 2}', f'{v["kw"]} of {v["edits"] * 2}']
        for k, v in by_type.items()
    ]
    summary.append(["All edits", len(rows), f'{run["correctChecks"]} of {run["totalChecks"]}', f'{run["keywordCorrectChecks"]} of {run["totalChecks"]}'])
    ws = sheet(workbook, "Summary", ["Edit type", "Edits", "AI checks correct", "Keyword-only checks correct"], summary, [28, 8, 20, 26])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    check_rows = []
    for r in rows:
        for phase, label in (("before", "Original"), ("after", "Modified")):
            correct = r[f"{phase}Correct"]
            check_rows.append([
                r["document"], r["agreementType"], r["provision"], EDIT_LABELS.get(r["editType"], r["editType"]),
                label, yes_no(r[f"expected{phase.title()}"]), yes_no(r[f"detected{phase.title()}"]),
                "Correct" if correct else "Miss",
                r[f"{phase}Confidence"],
                yes_no(r[f"{phase}QuotesEditedText"]),
                (r[f"{phase}Quote"] or "")[:4000] or None,
                yes_no(r[f"keyword{phase.title()}Correct"]),
                diagnoses.get(f'{r["document"]}|{r["provision"]}|{phase}') if not correct else None,
                r["description"],
            ])
    ws = sheet(
        workbook, "Checks",
        ["Document", "Agreement type", "Provision", "Edit", "Version", "Expected present", "AI flagged",
         "AI result", "Confidence", "AI quote touches edited text", "AI quote", "Keyword-only correct",
         "Diagnosis of miss", "Edit made"],
        check_rows, [26, 20, 26, 22, 10, 10, 10, 10, 10, 14, 70, 12, 60, 60], wrap=(11, 13, 14),
    )
    for row in ws.iter_rows(min_row=2):
        row[7].fill = OK_FILL if row[7].value == "Correct" else MISS_FILL

    notes = [
        ["What was run", f'Calder Automatic (AI) path, model {run["model"]}, {run["runAt"][:10]}: {run["documents"]} Bonterms/Common Paper agreements, each in its original and deliberately modified form ({run["aiCalls"]} AI calls, {run["aiFallbacks"]} fallbacks).'],
        ["Answer key", "evaluation/independent-check/labels.json. We made each edit, so we know whether the edited provision should be present before and after it."],
        ["Pipeline", "Same as the app: each file submitted as a TXT upload with its own agreement type's playbook; api/analyze.js; quotes not found verbatim in the document are dropped; confidence 90 if keyword matching agrees, else 75."],
        ["Unit", "Two checks per edit: is the edited provision flagged on the original, and on the modified agreement?"],
        ["Quote touches edited text", "Whether the AI's quote shares text with the lines the edit removed (original) or added (modified). Shows whether a correct flag points at the clause that changed."],
        ["Keyword-only", "The deterministic path's committed results for the same checks (evaluation/independent-check/results/), for comparison."],
        ["Test-format caveat", "The agreements are Markdown with **bold** and <span> markup. The AI quotes the wording without the markup, so the exact-quote check drops some correct findings. A PDF or plain-text upload would not carry this markup."],
    ]
    ws = workbook.create_sheet("Method")
    for row in notes:
        ws.append(row)
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 120
    for cell in ws["A"]:
        cell.font = Font(bold=True)
    for cell in ws["B"]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    args.out.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.out)
    print(f"Wrote {args.out}")
    for row in summary:
        print(f"  {row[0]:26} edits {row[1]:2}  AI {row[2]:9}  keyword {row[3]}")


if __name__ == "__main__":
    main()
