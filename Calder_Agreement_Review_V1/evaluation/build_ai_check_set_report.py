"""Score an AI check-set run against CUAD's expert labels and write an Excel report.

Sections 6.1/6.2 of the brief: for each provision category, how often Calder
finds what the labels say is there, and how often its flags are wrong. No
single overall accuracy figure is computed anywhere in this report.

Usage:
  python3 evaluation/build_ai_check_set_report.py \
    --results evaluation/ai-check-set/results.json \
    --dataset evaluation/data/CUAD_v1/CUAD_v1.json \
    --out "Calder AI Check Set.xlsx" \
    --sql evaluation/ai-check-set/load_results.sql

--sql also writes a file to paste into the Supabase SQL Editor, which loads
the run into the evaluation tables shown on the app's AI evaluation page.

Requires openpyxl.
"""

from __future__ import annotations

import argparse
import json
import pathlib
import re
from difflib import SequenceMatcher

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parents[1]
TAXONOMY_PATH = ROOT / "src" / "data" / "cuad-calder-taxonomy.json"
LOW_SUPPORT = 5


def normalize_category(value: str) -> str:  # same rule as evaluate_cuad.py
    return re.sub(r"[^a-z0-9]", "", value.lower())


def category_from_question(qa: dict) -> str:  # same rule as evaluate_cuad.py
    identifier = qa.get("id", "")
    if "__" in identifier:
        return identifier.rsplit("__", 1)[1]
    match = re.search(r'related to "([^"]+)"', qa.get("question", ""))
    return match.group(1) if match else qa.get("question", "")


def norm_text(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip().lower()


def quotes_overlap(quote: str, clauses: list[str]) -> bool:
    """True when the Calder quote and an expert-labeled clause share a substantial run of text."""
    a = norm_text(quote)
    for clause in clauses:
        b = norm_text(clause)
        if not a or not b:
            continue
        if a in b or b in a:
            return True
        match = SequenceMatcher(None, a, b, autojunk=False).find_longest_match(0, len(a), 0, len(b))
        if match.size >= min(40, max(15, min(len(a), len(b)) // 2)):
            return True
    return False


def expert_labels(dataset_path: pathlib.Path) -> dict[str, dict[str, list[str]]]:
    """title -> normalized CUAD category -> expert answer spans (empty list = labeled absent)."""
    data = json.loads(dataset_path.read_text(encoding="utf-8"))["data"]
    labels: dict[str, dict[str, list[str]]] = {}
    for document in data:
        categories: dict[str, list[str]] = {}
        for paragraph in document.get("paragraphs", []):
            for qa in paragraph.get("qas", []):
                key = normalize_category(category_from_question(qa))
                spans = categories.setdefault(key, [])
                if not qa.get("is_impossible", False):
                    spans.extend(answer["text"] for answer in qa.get("answers", []) if answer.get("text"))
        labels[document["title"]] = categories
    return labels


def sql_literal(value) -> str:
    if value is None:
        return "null"
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (int, float)):
        return str(value)
    return "'" + str(value).replace("\x00", "").replace("'", "''") + "'"


OUTCOME_KEYS = {"Found": "found", "Missed": "missed", "Wrong flag": "wrong_flag", "Correctly silent": "correctly_silent"}


def write_sql(run: dict, detail_rows: list[list], path: pathlib.Path) -> str:
    """SQL for the Supabase SQL Editor: replaces this run's rows in one transaction."""
    run_id = f'{run["runAt"][:10]}-{run["model"]}-{len(run["results"])}-contracts'
    fallbacks = sum(1 for r in run["results"] if not r["aiSucceeded"])
    label = f'{len(run["results"])} CUAD contracts, Automatic (AI), {run["model"]}, {run["runAt"][:10]}'
    notes = (
        f'Agreement type "{run["agreementType"]}" (all ten provisions). Answer key: CUAD v1 expert labels. '
        "Unit: one contract x one category; flagged = at least one finding for that category."
    )
    lines = [
        "-- Calder AI evaluation results. Paste into Supabase SQL Editor and run.",
        "-- Requires the evaluation tables from supabase/schema.sql.",
        "begin;",
        f"delete from public.evaluation_runs where id = {sql_literal(run_id)};",
        "insert into public.evaluation_runs (id,label,run_at,model,agreement_type,contracts,ai_calls,ai_fallbacks,notes) values ("
        + ",".join(sql_literal(v) for v in [
            run_id, label, run["runAt"], run["model"], run["agreementType"], len(run["results"]),
            sum(r["aiCalls"] for r in run["results"]), fallbacks, notes,
        ]) + ");",
    ]
    values = []
    for row in detail_rows:
        contract, provision, expert, flagged, outcome, confidence, quote, overlap, clauses, kw, method = row
        values.append("(" + ",".join(sql_literal(v) for v in [
            run_id, contract, provision, expert == "Present", flagged == "Yes", OUTCOME_KEYS[outcome],
            confidence, quote[:2000] if quote else None,
            None if overlap == "-" else overlap == "Yes",
            clauses[:2000] if clauses else None, kw == "Yes", method,
        ]) + ")")
    for start in range(0, len(values), 100):
        lines.append(
            "insert into public.evaluation_results (run_id,contract,provision,expert_present,calder_flagged,outcome,confidence,calder_quote,quote_overlaps,expert_clause,keyword_flagged,method_used) values\n"
            + ",\n".join(values[start:start + 100]) + ";"
        )
    lines.append("commit;")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return run_id


def pct(numerator: int, denominator: int) -> str:
    return f"{numerator / denominator:.0%}" if denominator else "n/a"


def pattern(found: int, labeled: int, wrong: int, flagged: int) -> str:
    if labeled < LOW_SUPPORT:
        return "Too few labeled examples to judge"
    recall = found / labeled
    false_rate = wrong / flagged if flagged else 0.0
    if recall >= 0.8 and false_rate <= 0.2:
        return "Reliable"
    if recall >= 0.8:
        return "Finds most, but over-flags"
    if false_rate <= 0.2:
        return "Precise, but misses often"
    return "Misses often and over-flags"


HEADER_FILL = PatternFill("solid", fgColor="1F3A4D")
HEADER_FONT = Font(bold=True, color="FFFFFF")
OUTCOME_FILL = {
    "Found": PatternFill("solid", fgColor="DCEFE3"),
    "Missed": PatternFill("solid", fgColor="F8DADA"),
    "Wrong flag": PatternFill("solid", fgColor="FBE7C6"),
}


def write_sheet(workbook, title, headers, rows, widths, wrap_columns=()):
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in rows:
        sheet.append(row)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for column in wrap_columns:
        for cell in sheet[get_column_letter(column)][1:]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions
    return sheet


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--results", required=True, type=pathlib.Path)
    parser.add_argument("--dataset", required=True, type=pathlib.Path)
    parser.add_argument("--out", required=True, type=pathlib.Path)
    parser.add_argument("--sql", type=pathlib.Path, help="Also write a SQL file to load the results into Supabase")
    args = parser.parse_args()

    run = json.loads(args.results.read_text(encoding="utf-8"))
    taxonomy = {e["calderProvision"]: e for e in json.loads(TAXONOMY_PATH.read_text(encoding="utf-8"))["provisions"]}
    labels = expert_labels(args.dataset)
    provisions = run["provisions"]

    detail_rows, per_category = [], {p: dict(labeled=0, flagged=0, found=0, missed=0, wrong=0, overlap=0, kw_flagged=0, kw_found=0, kw_wrong=0) for p in provisions}
    for result in run["results"]:
        doc_labels = labels[result["contract"]]
        for provision in provisions:
            cuad_keys = [normalize_category(c) for c in taxonomy[provision]["cuadCategories"]]
            spans = [s for k in cuad_keys for s in doc_labels.get(k, [])]
            present = bool(spans)
            hits = [f for f in result["findings"] if f["provision"] == provision]
            kw_hit = any(f["provision"] == provision for f in result["deterministicFindings"])
            flagged = bool(hits)
            outcome = ("Found" if present else "Wrong flag") if flagged else ("Missed" if present else "Correctly silent")
            overlap = "-" if not (flagged and present) else ("Yes" if any(quotes_overlap(h["sourceText"], spans) for h in hits) else "No")
            stats = per_category[provision]
            stats["labeled"] += present
            stats["flagged"] += flagged
            stats["found"] += flagged and present
            stats["missed"] += present and not flagged
            stats["wrong"] += flagged and not present
            stats["overlap"] += overlap == "Yes"
            stats["kw_flagged"] += kw_hit
            stats["kw_found"] += kw_hit and present
            stats["kw_wrong"] += kw_hit and not present
            detail_rows.append([
                result["contract"], provision, "Present" if present else "Absent",
                "Yes" if flagged else "No", outcome,
                max((h["confidence"] for h in hits), default=None),
                "\n---\n".join(h["sourceText"] for h in hits)[:4000] or None,
                overlap,
                "\n---\n".join(spans[:3])[:4000] or None,
                "Yes" if kw_hit else "No",
                result["methodUsed"],
            ])

    workbook = Workbook()
    workbook.remove(workbook.active)

    summary_rows = []
    for provision in provisions:
        s = per_category[provision]
        summary_rows.append([
            provision, s["labeled"], s["found"], s["missed"],
            f'{s["found"]} of {s["labeled"]} ({pct(s["found"], s["labeled"])})',
            s["flagged"], s["wrong"],
            f'{s["wrong"]} of {s["flagged"]} ({pct(s["wrong"], s["flagged"])})',
            f'{s["overlap"]} of {s["found"]}' if s["found"] else "n/a",
            pattern(s["found"], s["labeled"], s["wrong"], s["flagged"]),
            f'{s["kw_found"]} of {s["labeled"]} ({pct(s["kw_found"], s["labeled"])})',
            f'{s["kw_wrong"]} of {s["kw_flagged"]} ({pct(s["kw_wrong"], s["kw_flagged"])})',
            None,
        ])
    write_sheet(
        workbook, "Per category",
        ["Provision category", "Contracts labeled present", "Found", "Missed",
         "How often it finds what is there", "Contracts Calder flagged", "Wrong flags",
         "How often its flags are wrong", "Found flags quoting the labeled clause",
         "Suggested reading", "Keyword-only: finds what is there", "Keyword-only: flags wrong",
         "Team explanation"],
        summary_rows, [28, 12, 9, 9, 22, 12, 10, 22, 18, 28, 22, 20, 60], wrap_columns=(13,),
    )

    detail = write_sheet(
        workbook, "Contract x category",
        ["Contract", "Provision category", "Expert label", "Calder flagged", "Outcome",
         "Confidence", "Calder quote", "Quote overlaps labeled clause", "Expert-labeled clause(s)",
         "Keyword-only flagged", "Method used"],
        detail_rows, [44, 26, 12, 10, 16, 11, 70, 14, 70, 12, 22], wrap_columns=(7, 9),
    )
    for row in detail.iter_rows(min_row=2):
        fill = OUTCOME_FILL.get(row[4].value)
        if fill:
            row[4].fill = fill

    write_sheet(
        workbook, "To spot-check",
        ["Contract", "Provision category", "Expert label", "Calder flagged", "Outcome",
         "Confidence", "Calder quote", "Quote overlaps labeled clause", "Expert-labeled clause(s)",
         "Keyword-only flagged", "Method used", "Your note (6.1)"],
        [r + [None] for r in detail_rows if r[4] in ("Missed", "Wrong flag") or r[7] == "No"],
        [44, 26, 12, 10, 16, 11, 70, 14, 70, 12, 22, 50], wrap_columns=(7, 9, 12),
    )

    write_sheet(
        workbook, "Contracts",
        ["Contract", "Pages", "Characters", "AI sections", "AI calls", "Method used",
         "AI error", "AI time (s)", "AI quotes dropped (not found in text)"],
        [[r["contract"], r["pages"], r["characters"], r["sections"], r["aiCalls"], r["methodUsed"],
          r["aiError"], round(r["aiMs"] / 1000, 1), len(r["unverifiedQuotes"])] for r in run["results"]],
        [60, 8, 12, 11, 9, 26, 30, 11, 16],
    )

    fallbacks = sum(1 for r in run["results"] if not r["aiSucceeded"])
    notes = [
        ["What was run", f'Calder Automatic (AI) path, model {run["model"]}, on {len(run["results"])} CUAD contracts, {run["runAt"][:10]}.'],
        ["Pipeline", "Same as the app: PDF.js text extraction; split into sections of up to 120,000 characters; api/analyze.js per section; quotes not found in the document dropped; confidence 90 if keyword matching found the same category, else 75."],
        ["Agreement type", f'Every contract was submitted as "{run["agreementType"]}", which carries all ten provisions.'],
        ["AI fallbacks", f"{fallbacks} contract(s) fell back to keyword matching because an AI call failed; their rows say so in Method used."],
        ["Answer key", "CUAD v1 expert labels. A category is labeled present when CUAD has at least one expert span for its mapped CUAD category. No labels were created for this check."],
        ["Unit", "One contract x one category. A contract counts as flagged for a category when Calder produced at least one finding for it."],
        ["Finds what is there", "Found / contracts labeled present (recall)."],
        ["Flags wrong", "Wrong flags / contracts Calder flagged (the share of flags with no expert label for that category)."],
        ["Quote overlap", "For Found rows only: whether Calder's quoted text shares a substantial run of text with an expert-labeled clause. A Found flag that quotes a different clause is right about presence but points the reviewer elsewhere."],
        ["Keyword-only columns", "The same contracts scored with Calder's deterministic keyword matching, for comparison. No AI involved."],
        ["Suggested reading", f"Rule of thumb only: reliable = finds at least 80% and at most 20% of flags wrong. Categories with fewer than {LOW_SUPPORT} labeled contracts are not judged."],
        ["Not reported", "No single overall accuracy figure, by design (brief Section 6.2)."],
        ["Spot-check (6.1)", "The 'To spot-check' sheet lists every miss, wrong flag, and found flag quoting a different clause, with the expert clause beside Calder's quote, for checking against the Labeling Handbook."],
    ]
    method = workbook.create_sheet("Method")
    for row in notes:
        method.append(row)
    method.column_dimensions["A"].width = 24
    method.column_dimensions["B"].width = 120
    for cell in method["A"]:
        cell.font = Font(bold=True)
    for cell in method["B"]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    args.out.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.out)
    print(f"Wrote {args.out}")
    if args.sql:
        run_id = write_sql(run, detail_rows, args.sql)
        print(f"Wrote {args.sql} (run {run_id}, {args.sql.stat().st_size // 1024} KB)")
    for row in summary_rows:
        print(f"  {row[0]:28} finds {row[4]:18} flags wrong {row[7]:18} {row[9]}")


if __name__ == "__main__":
    main()
